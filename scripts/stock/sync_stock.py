#!/usr/bin/env python3
"""Sync de stock · Excel de fábrica (FTPS) → inventario de Shopify.

Port de la lógica de `cartujasync` (docs/cartujasync-especificacion.docx.md):

  1. Baja `MASTER PRODUCCION.xlsm` e `Inventarios_NOBORRAR.xlsm` del FTP de fábrica
     (FTPS, FileZilla Server). Mismas pestañas Producto/Bizcocho (Inventarios trae ~46
     referencias más; si coinciden manda el MASTER) y Inventarios añade `RelPiezas`,
     la composición de cada conjunto según fábrica.
  2. Por referencia LCS: vendible = AL1-Primera − Pte. Servir − pedidos web que
     aún no están en Navision (pedidos de Shopify sin el tag `navision`).
     Bizcocho por IdForma+IdPieza (pestaña Bizcocho).
  3. Cada variante se reduce a una «receta» de referencias del Excel:
       simple        SKU = referencia                                 {ref: 1}
       conjunto      pieza en RelPiezas (azucarero = cuerpo + tapa, taza con platillo,
                     juego de boles = 6 × bol…); IdForma vacío = la forma del SKU
       pack          `custom.pack` → metaobjects pieza_pack (producto × cantidad)
     y la manda el componente más restrictivo.
  4. Caso 1 stock (7 días) · 2 bizcocho (15) · 3 fabricar desde cero (21) · 4 no disponible.
     Como el módulo de PrestaShop: si la pieza es fabricable (tiene línea de bizcocho) se
     vende aunque no haya stock (política CONTINUE, «bajo pedido»); si no, DENY.
     Inventario = stock + bizcocho (cifra real). TOPE_CASO3 > 0 limita la fase 3 a ese
     total (DENY con inventario = tope); 0 = sin tope, igual que la tienda antigua.
     El desglose va al metafield `cartuja.stock` para que el tema pinte el plazo.

Tag `fin-de-existencias`: no se vuelven a fabricar, así que sólo cuenta el stock
terminado (sin bizcocho ni fabricación: caso 1 o 4). Los que no están en el Excel no
se tocan (stock manual). Tampoco se tocan las variantes sin SKU o sin receta (salen
en el informe).

    python3 sync_stock.py                    # ensayo: informe, no escribe
    python3 sync_stock.py --apply            # escribe sólo lo que cambia
    python3 sync_stock.py --local carpeta/   # sin FTP, con los dos .xlsm de esa carpeta

Credenciales en entorno o en `gropius/.env`: FTP_HOST FTP_PORT FTP_USER FTP_PASS
FTP_DIR SHOPIFY_STORE SHOPIFY_CLIENT_ID SHOPIFY_CLIENT_SECRET.
"""

import argparse
import datetime as dt
import ftplib
import json
import os
import sys
import tempfile
import time
import urllib.parse
import urllib.request
import uuid
from pathlib import Path

import openpyxl

API = "2026-10"
LOCATION = "Almacén Sevilla"
MASTER = "MASTER PRODUCCION.xlsm"
INVENTARIOS = "Inventarios_NOBORRAR.xlsm"
FIN_TAG = "fin-de-existencias"
NAVISION_TAG = "navision"
TOPE_CASO3 = 0      # 0 = sin tope (como CARTUJASYNC_DELIVERY3_MAX_QTY en PrestaShop); N = máx. unidades en fase 3
MAX_AGE_H = 36      # Excel más viejo → salta la pasada sin escribir ni fallar (fábrica no lo sube en fin de semana)
DIAS = {1: 7, 2: 15, 3: 21}     # 21: lo que mostraba la tienda antigua y dicen las políticas («15–21»)


# ---- entorno ------------------------------------------------------------------
def load_env():
    env = dict(os.environ)
    for p in (Path(__file__).resolve().parents[3] / ".env", Path.cwd() / ".env"):
        if p.exists():
            for line in p.read_text().splitlines():
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, v = line.split("=", 1)
                    env.setdefault(k.strip(), v.strip().strip("'\""))
            break
    return env


# ---- FTP ----------------------------------------------------------------------
class ReuseTLS(ftplib.FTP_TLS):
    """FileZilla Server exige reutilizar la sesión TLS del control en el canal de datos."""

    def ntransfercmd(self, cmd, rest=None):
        conn, size = ftplib.FTP.ntransfercmd(self, cmd, rest)
        if self._prot_p:
            conn = self.context.wrap_socket(conn, server_hostname=self.host, session=self.sock.session)
        return conn, size


def download_excel(env, carpeta, intentos=3):
    # ponytail: reintento simple; el FTP de fábrica a veces no abre el puerto pasivo (timeout)
    for i in range(intentos):
        try:
            return _download(env, carpeta)
        except (OSError, ftplib.Error) as e:
            print("  FTP intento %d/%d: %s" % (i + 1, intentos, e))
            if i + 1 == intentos:
                raise
            time.sleep(30)


def _download(env, carpeta):
    ftp = ReuseTLS(timeout=30)
    ftp.connect(env["FTP_HOST"], int(env.get("FTP_PORT") or 21))
    ftp.login(env["FTP_USER"], env["FTP_PASS"])
    ftp.prot_p()
    if env.get("FTP_DIR"):
        ftp.cwd(env["FTP_DIR"])
    for name in (MASTER, INVENTARIOS):
        with open(os.path.join(carpeta, name), "wb") as fh:
            ftp.retrbinary("RETR " + name, fh.write)
    ftp.quit()


# ---- Excel --------------------------------------------------------------------
def read_excel(path):
    """→ (refs {ref: {stock, pendiente, forma, pieza}}, bizcocho {(forma, pieza): n},
    actualizado, conjuntos {pieza: [(pieza, forma|"", n)]} si hay pestaña RelPiezas)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # «Ultima Actualización» en la hoja de menú (MENU!B3→B5 en MASTER, Menu!G3→G5 en Inventarios)
    menu = wb[next(n for n in wb.sheetnames if n.upper() == "MENU")]
    grid = [list(r) for r in menu.iter_rows(max_row=10, values_only=True)]
    updated = next((grid[i + 2][j] for i, row in enumerate(grid[:-2]) for j, c in enumerate(row)
                    if isinstance(c, str) and "ltima Actualizaci" in c), None)
    if not isinstance(updated, dt.datetime):
        sys.exit("%s: no encuentro la fecha de «Ultima Actualización» (%r)" % (os.path.basename(path), updated))

    def sheet(name, cols):
        rows = wb[name].iter_rows(values_only=True)
        head = [str(c).strip() if c is not None else "" for c in next(rows)]
        missing = [c for c in cols if c not in head]
        if missing:
            sys.exit("Falta(n) columna(s) %s en la pestaña %s" % (missing, name))
        idx = [head.index(c) for c in cols]
        for r in rows:
            if r[0]:
                yield [r[i] for i in idx]

    refs = {}
    for ref, forma, pieza, stock, pend in sheet("Producto", ["Referencia LCS", "IdForma", "IdPieza", "AL1-Primera", "Pte. Servir"]):
        # la misma referencia puede venir en varias filas: se suman
        r = refs.setdefault(str(ref).strip().upper(), {"stock": 0, "pendiente": 0, "forma": str(forma).strip(), "pieza": str(pieza).strip()})
        r["stock"] += int(stock or 0)
        r["pendiente"] += int(pend or 0)
    biz = {}
    for forma, pieza, n in sheet("Bizcocho", ["IdForma", "IdPieza", "Stock bizcocho"]):
        k = (str(forma).strip(), str(pieza).strip())
        biz[k] = biz.get(k, 0) + int(n or 0)
    conjuntos = {}
    if "RelPiezas" in wb.sheetnames:
        for conj, pieza, forma, n in sheet("RelPiezas", ["IdConjunto", "IdPieza", "IdForma", "NumPiezas"]):
            forma = "" if forma in (None, "", "None") else str(forma).strip()
            conjuntos.setdefault(str(conj).strip(), []).append((str(pieza).strip(), forma, int(n or 1)))
    return refs, biz, updated, conjuntos


# ---- Shopify ------------------------------------------------------------------
class Shopify:
    def __init__(self, env):
        self.store = env["SHOPIFY_STORE"]
        body = urllib.parse.urlencode({"grant_type": "client_credentials", "client_id": env["SHOPIFY_CLIENT_ID"],
                                       "client_secret": env["SHOPIFY_CLIENT_SECRET"]}).encode()
        self.token = json.load(urllib.request.urlopen("https://%s/admin/oauth/access_token" % self.store, body))["access_token"]

    def gql(self, query, variables=None):
        req = urllib.request.Request("https://%s/admin/api/%s/graphql.json" % (self.store, API),
                                     json.dumps({"query": query, "variables": variables or {}}).encode(),
                                     {"Content-Type": "application/json", "X-Shopify-Access-Token": self.token})
        res = json.load(urllib.request.urlopen(req, timeout=60))
        for wait in (2, 5, 10, 20):   # cubo de coste agotado: esperar y repetir
            if not any(e.get("extensions", {}).get("code") == "THROTTLED" for e in res.get("errors") or []):
                break
            time.sleep(wait)
            res = json.load(urllib.request.urlopen(req, timeout=60))
        if res.get("errors"):
            raise RuntimeError(json.dumps(res["errors"], ensure_ascii=False))
        return res["data"]

    def paginate(self, query, root, variables=None):
        cursor = None
        while True:
            page = self.gql(query, dict(variables or {}, c=cursor))[root]
            yield from page["nodes"]
            if not page["pageInfo"]["hasNextPage"]:
                return
            cursor = page["pageInfo"]["endCursor"]


Q_VARIANTS = """query($c: String) { productVariants(first: 250, after: $c) { pageInfo { hasNextPage endCursor } nodes {
  id sku title inventoryPolicy
  product { id title status tags }
  inventoryItem { id tracked inventoryLevels(first: 10) { nodes { location { id name } quantities(names: ["available"]) { quantity } } } }
  pack: metafield(namespace: "custom", key: "pack") { value }
  stock: metafield(namespace: "cartuja", key: "stock") { value }
} } }"""
Q_NODES = "query($ids: [ID!]!) { nodes(ids: $ids) { ... on Metaobject { id fields { key value } } } }"
Q_ORDERS = """query($c: String, $q: String) { orders(first: 100, after: $c, query: $q) { pageInfo { hasNextPage endCursor } nodes {
  name lineItems(first: 100) { nodes { currentQuantity variant { id } } }
} } }"""
Q_LOCATIONS = "{ locations(first: 20) { nodes { id name } } }"

# 2026-10: las mutaciones de inventario exigen @idempotent (una clave por llamada)
M_SET = """mutation($input: InventorySetQuantitiesInput!, $key: String!) { inventorySetQuantities(input: $input) @idempotent(key: $key) {
  userErrors { field message code } } }"""
M_ACTIVATE = """mutation($item: ID!, $loc: ID!, $key: String!) { inventoryActivate(inventoryItemId: $item, locationId: $loc) @idempotent(key: $key) {
  userErrors { field message } } }"""
M_ITEM = """mutation($id: ID!, $input: InventoryItemInput!) { inventoryItemUpdate(id: $id, input: $input) {
  userErrors { field message } } }"""
M_VARIANTS = """mutation($pid: ID!, $v: [ProductVariantsBulkInput!]!) { productVariantsBulkUpdate(productId: $pid, variants: $v) {
  userErrors { field message } } }"""
M_META = """mutation($m: [MetafieldsSetInput!]!) { metafieldsSet(metafields: $m) { userErrors { field message } } }"""


# ---- cálculo ------------------------------------------------------------------
def recipe(sku, refs, conjuntos, nivel=0):
    """SKU → {ref: cantidad} o None si no hay forma de sacarlo del Excel.

    SKU = forma(2) + decorado(3) + pieza(4). Si no es una referencia del Excel y la pieza
    es un conjunto de RelPiezas, se compone con sus piezas (misma forma y decorado salvo
    que RelPiezas fije otra forma: la taza de desayuno es siempre forma 07)."""
    sku = (sku or "").strip().upper()
    if not sku:
        return None
    if sku in refs:
        return {sku: 1}
    if len(sku) != 9 or sku[5:] not in conjuntos or nivel > 3:
        return None
    total = {}
    for pieza, forma, n in conjuntos[sku[5:]]:
        sub = recipe((forma or sku[:2]) + sku[2:5] + pieza, refs, conjuntos, nivel + 1)
        if not sub:
            return None
        for r, q in sub.items():
            total[r] = total.get(r, 0) + q * n
    return total


def disponibilidad(rec, avail, biz):
    """Receta → (caso, stock_packs, bizcocho_packs, inventario, política, fabricable).
    El componente más restrictivo manda."""
    stock = min(avail[r] // q for r, q in rec.items())
    total = min((avail[r] + (biz[r] or 0)) // q for r, q in rec.items())
    fabricable = all(biz[r] is not None for r in rec)  # todas las piezas tienen línea de bizcocho
    if stock >= 1:
        caso = 1
    elif total >= 1:
        caso = 2
    elif fabricable:
        caso = 3
    else:
        caso = 4
    if fabricable and TOPE_CASO3 == 0:
        return caso, stock, total - stock, total, "CONTINUE", True      # bajo pedido sin límite
    inventario = max(total, TOPE_CASO3) if fabricable else total
    return caso, stock, total - stock, inventario, "DENY", fabricable


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--apply", action="store_true", help="escribe en Shopify (sin esto: ensayo)")
    ap.add_argument("--local", help="carpeta con los dos .xlsm (en vez de bajarlos del FTP)")
    ap.add_argument("--report", default="informe-sync-stock.json")
    args = ap.parse_args()
    env = load_env()

    carpeta = args.local
    if not carpeta:
        carpeta = tempfile.mkdtemp()
        download_excel(env, carpeta)
    refs, biz_by_piece, updated, _ = read_excel(os.path.join(carpeta, MASTER))
    refs_inv, biz_inv, updated_inv, conjuntos = read_excel(os.path.join(carpeta, INVENTARIOS))
    extra = len(refs_inv.keys() - refs.keys())
    refs, biz_by_piece = {**refs_inv, **refs}, {**biz_inv, **biz_by_piece}   # si coinciden, manda el MASTER
    age_h = (dt.datetime.now() - updated).total_seconds() / 3600
    print("Excel: %d referencias (+%d de Inventarios) · %d conjuntos · actualizado %s (hace %.1f h; Inventarios %s)"
          % (len(refs), extra, len(conjuntos), updated, age_h, updated_inv))
    if age_h > MAX_AGE_H:
        print("· El Excel tiene más de %d h (sin cambios en fábrica): no se toca nada." % MAX_AGE_H)
        sys.exit(0)

    shop = Shopify(env)
    loc = next((l["id"] for l in shop.gql(Q_LOCATIONS)["locations"]["nodes"] if l["name"] == LOCATION), None)
    if not loc:
        sys.exit("✗ No existe la ubicación %r" % LOCATION)
    variants = list(shop.paginate(Q_VARIANTS, "productVariants"))
    by_id = {v["id"]: v for v in variants}

    # packs: custom.pack → [pieza_pack{producto, cantidad}] → receta sumando las de sus piezas
    pack_ids = sorted({g for v in variants if v["pack"] for g in json.loads(v["pack"]["value"])})
    piezas = {}
    for i in range(0, len(pack_ids), 100):
        for n in shop.gql(Q_NODES, {"ids": pack_ids[i:i + 100]})["nodes"]:
            if n:
                piezas[n["id"]] = {f["key"]: f["value"] for f in n["fields"]}
    first_variant = {}
    for v in variants:
        first_variant.setdefault(v["product"]["id"], v)

    def receta_variante(v):
        if not v["pack"]:
            return recipe(v["sku"], refs, conjuntos), None
        total = {}
        for g in json.loads(v["pack"]["value"]):
            p = piezas.get(g)
            comp = p and first_variant.get(p.get("producto"))
            sub = comp and recipe(comp["sku"], refs, conjuntos)
            if not sub:
                return None, "pieza de pack sin resolver (%s)" % (g if not p else p.get("producto"))
            for r, q in sub.items():
                total[r] = total.get(r, 0) + q * int(p.get("cantidad") or 1)
        return total, None

    recetas, sin_receta, excluidas, fin = {}, [], [], set()
    for v in variants:
        rec, why = receta_variante(v)
        if FIN_TAG in v["product"]["tags"]:
            if not rec:
                excluidas.append(v["id"])   # fin de existencias fuera del Excel: stock manual
                continue
            fin.add(v["id"])
        if rec:
            recetas[v["id"]] = rec
        else:
            sin_receta.append({"sku": v["sku"], "producto": v["product"]["title"], "variante": v["title"],
                               "motivo": why or ("sin SKU" if not (v["sku"] or "").strip() else "SKU no está en el Excel")})

    # pedidos web que Navision aún no conoce: se descuentan (ponytail: read_orders sólo ve 60 días)
    desde = (dt.date.today() - dt.timedelta(days=60)).isoformat()
    q = "created_at:>=%s -tag:%s -status:cancelled" % (desde, NAVISION_TAG)
    web, pedidos = {}, []
    for o in shop.paginate(Q_ORDERS, "orders", {"q": q}):
        pedidos.append(o["name"])
        for li in o["lineItems"]["nodes"]:
            rec = li["variant"] and (recetas.get(li["variant"]["id"]) or receta_variante(by_id.get(li["variant"]["id"], {"pack": None, "sku": ""}))[0])
            for r, n in (rec or {}).items():
                web[r] = web.get(r, 0) + n * li["currentQuantity"]   # sin lo reembolsado o quitado

    avail = {r: max(0, d["stock"] - d["pendiente"] - web.get(r, 0)) for r, d in refs.items()}
    biz = {r: biz_by_piece.get((d["forma"], d["pieza"])) for r, d in refs.items()}

    cambios, resumen = [], {1: 0, 2: 0, 3: 0, 4: 0}
    for v in variants:
        rec = recetas.get(v["id"])
        if not rec:
            continue
        # fin de existencias: no se fabrican más → ni bizcocho ni fabricación
        caso, stock, bizc, pedible, politica, fab = disponibilidad(rec, avail, {r: None for r in rec} if v["id"] in fin else biz)
        resumen[caso] += 1
        levels = {l["location"]["id"]: l["quantities"][0]["quantity"] for l in v["inventoryItem"]["inventoryLevels"]["nodes"]}
        # lo lee snippets/plazo-entrega.liquid: caso según cantidad = stock → bizcocho → fabricable
        # (fab) hasta el tope (0 = sin tope)
        meta = json.dumps({"caso": caso, "stock": stock, "bizcocho": bizc, "fab": fab,
                           "tope": TOPE_CASO3 if fab else 0,
                           "dias": [DIAS[1], DIAS[2], DIAS[3]], "excel": updated.isoformat()}, separators=(",", ":"))
        old_meta = json.loads(v["stock"]["value"]) if v["stock"] else {}
        c = {
            "id": v["id"], "sku": v["sku"], "producto": v["product"]["title"], "variante": v["title"],
            "item": v["inventoryItem"]["id"], "pid": v["product"]["id"], "caso": caso,
            "antes": levels.get(loc), "despues": pedible,
            "otras": {k: n for k, n in levels.items() if k != loc and n},   # stock en otras ubicaciones → a 0
            "activar": loc not in levels,
            "tracked": not v["inventoryItem"]["tracked"],
            "politica": politica if v["inventoryPolicy"] != politica else None,
            "meta": meta if {k: x for k, x in old_meta.items() if k != "excel"} != {k: x for k, x in json.loads(meta).items() if k != "excel"} else None,
        }
        if c["antes"] != pedible or c["otras"] or c["activar"] or c["tracked"] or c["politica"] or c["meta"]:
            cambios.append(c)

    print("Variantes: %d · con receta %d (de ellas %s: %d) · sin receta %d · %s fuera del Excel (manual): %d"
          % (len(variants), len(recetas), FIN_TAG, len(fin), len(sin_receta), FIN_TAG, len(excluidas)))
    print("Casos: " + " · ".join("%d→%d" % kv for kv in resumen.items()))
    print("Pedidos web sin tag %r (60 días): %d" % (NAVISION_TAG, len(pedidos)))   # sin números: el log del workflow es público
    print("Cambios: %d (cantidad %d · activar en %s %d · tracked %d · política %d · metafield %d · otras ubic. %d)" % (
        len(cambios), sum(c["antes"] != c["despues"] for c in cambios), LOCATION, sum(c["activar"] for c in cambios),
        sum(c["tracked"] for c in cambios), sum(bool(c["politica"]) for c in cambios), sum(bool(c["meta"]) for c in cambios),
        sum(bool(c["otras"]) for c in cambios)))

    fallos = aplicar(shop, loc, cambios) if args.apply else 0

    with open(args.report, "w") as fh:
        json.dump({"modo": "aplicado" if args.apply else "ensayo", "excel_actualizado": updated.isoformat(),
                   "pedidos_web_descontados": pedidos, "casos": resumen, "cambios": cambios,
                   "sin_receta": sin_receta}, fh, indent=1, ensure_ascii=False)
    print("Informe en %s" % args.report)
    if fallos:
        sys.exit("✗ %d escrituras con error: se reintentan en la siguiente pasada" % fallos)


def aplicar(shop, loc, cambios):
    """Escribe los cambios. Devuelve cuántas llamadas fallaron."""
    fallos = []

    def check(res, root, label):
        errs = res[root]["userErrors"]
        if errs:
            print("  ✗ %s: %s" % (label, json.dumps(errs, ensure_ascii=False)))
            fallos.append(label)
        return not errs

    for c in cambios:
        if c["tracked"]:
            check(shop.gql(M_ITEM, {"id": c["item"], "input": {"tracked": True}}), "inventoryItemUpdate", c["sku"])
        if c["activar"]:
            check(shop.gql(M_ACTIVATE, {"item": c["item"], "loc": loc, "key": str(uuid.uuid4())}), "inventoryActivate", c["sku"])
            c["antes"] = 0
        if c["politica"]:
            check(shop.gql(M_VARIANTS, {"pid": c["pid"], "v": [{"id": c["id"], "inventoryPolicy": c["politica"]}]}),
                  "productVariantsBulkUpdate", c["sku"])
        if c["meta"]:
            check(shop.gql(M_META, {"m": [{"ownerId": c["id"], "namespace": "cartuja", "key": "stock", "type": "json",
                                           "value": c["meta"]}]}), "metafieldsSet", c["sku"])

    # cantidades en lotes; changeFromQuantity hace que falle si alguien compró entre la lectura y la escritura
    qs = []
    for c in cambios:
        if c["antes"] != c["despues"]:
            qs.append({"inventoryItemId": c["item"], "locationId": loc, "quantity": c["despues"], "changeFromQuantity": c["antes"]})
        for other, n in c["otras"].items():
            qs.append({"inventoryItemId": c["item"], "locationId": other, "quantity": 0, "changeFromQuantity": n})
    for i in range(0, len(qs), 250):
        check(shop.gql(M_SET, {"key": str(uuid.uuid4()), "input": {"name": "available", "reason": "correction",
                                         "referenceDocumentUri": "cartuja-sync://excel", "quantities": qs[i:i + 250]}}),
              "inventorySetQuantities", "lote %d" % (i // 250 + 1))
    print("Aplicado: %d cambios de cantidad · %d llamadas con error" % (len(qs), len(fallos)))
    return len(fallos)


if __name__ == "__main__":
    main()
